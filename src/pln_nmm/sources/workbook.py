"""Read an NMM model-input workbook into plain dicts.

Deliberately thin: this only reads. Rule enforcement lives in
scripts/validate_workbook.py, and the builder refuses to run on a workbook that
has not passed it.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

HEADER_ROW = 2
DATA_START = 3


@dataclass
class Workbook:
    """Sheet name -> list of row dicts, keyed by bare header name."""

    sheets: dict[str, list[dict]]

    def rows(self, sheet: str) -> list[dict]:
        return self.sheets.get(sheet, [])

    def gi_skema_map(self) -> dict[str, str]:
        return {
            str(r["gi_id"]): str(r.get("skema_busbar") or "DOUBLE")
            for r in self.rows("01_GI")
            if r.get("gi_id")
        }

    def index_by(self, sheet: str, key: str) -> dict[str, dict]:
        return {str(r[key]): r for r in self.rows(sheet) if r.get(key)}


def _read_sheet(ws) -> list[dict]:
    headers = []
    for cell in ws[HEADER_ROW]:
        v = cell.value
        headers.append(v.replace(" *", "").strip() if isinstance(v, str) else None)

    rows: list[dict] = []
    for r in range(DATA_START, ws.max_row + 1):
        values: dict = {}
        empty = True
        for i, h in enumerate(headers, start=1):
            if not h:
                continue
            v = ws.cell(row=r, column=i).value
            if isinstance(v, str):
                v = v.strip() or None
            values[h] = v
            if v is not None:
                empty = False
        if not empty:
            values["_row"] = r
            rows.append(values)
    return rows


def read_workbook(path: str | Path) -> Workbook:
    wb = load_workbook(path, data_only=True)
    return Workbook(sheets={name: _read_sheet(wb[name]) for name in wb.sheetnames})
