"""Validate an NMM model-input workbook before it reaches the CIM generator.

Checks structure, controlled vocabularies, cross-sheet references, and the
data-quality rules that PLN's existing CIM files violate (zeros standing in for
unknown impedance, assumptions indistinguishable from facts, transformers with
identical HV/LV windings).

Usage:
    python scripts/validate_workbook.py NMM_Model_Input_TEMPLATE.xlsx
"""

from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from build_workbook_template import (
    BAY_TIPE,
    BUS_NORMAL,
    CONFIDENCE,
    JENIS_NILAI,
    PHI,
    SHEETS,
    SKEMA_BUSBAR,
    SUMBER,
    TEMPLATE_ID,
    TITIK_UKUR,
)

HEADER_ROW = 2
DATA_START = 3


class Report:
    """Collects findings. Errors block generation; warnings do not."""

    def __init__(self):
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.info: list[str] = []

    def error(self, sheet, row, msg):
        loc = f"{sheet}!{row}" if row else sheet
        self.errors.append(f"{loc}: {msg}")

    def warn(self, sheet, row, msg):
        loc = f"{sheet}!{row}" if row else sheet
        self.warnings.append(f"{loc}: {msg}")

    @property
    def ok(self):
        return not self.errors


def _read_sheet(ws):
    """Return (headers, rows) where rows are dicts keyed by bare header name."""
    headers = []
    for cell in ws[HEADER_ROW]:
        v = cell.value
        headers.append(v.replace(" *", "").strip() if isinstance(v, str) else None)

    rows = []
    for r in range(DATA_START, ws.max_row + 1):
        values = {}
        empty = True
        for i, h in enumerate(headers, start=1):
            if not h:
                continue
            v = ws.cell(row=r, column=i).value
            if isinstance(v, str):
                v = v.strip()
                if v == "":
                    v = None
            values[h] = v
            if v is not None:
                empty = False
        if not empty:
            values["_row"] = r
            rows.append(values)
    return headers, rows


def _spec_map():
    """Expected columns and required flags, from the template definition."""
    out = {}
    for fn in SHEETS:
        name, _note, cols, _rows = fn()
        out[name] = {hdr: required for hdr, _w, required, _v in cols}
    return out


def _check_vocab(rep, sheet, rows, column, allowed):
    for row in rows:
        v = row.get(column)
        if v is None:
            continue
        if str(v) not in allowed:
            rep.error(
                sheet,
                row["_row"],
                f"{column}='{v}' bukan nilai sah. Pilihan: {', '.join(allowed)}",
            )


def _check_required(rep, sheet, rows, spec):
    for row in rows:
        for col, required in spec.items():
            if required and row.get(col) is None:
                rep.error(sheet, row["_row"], f"kolom wajib '{col}' kosong")


def _check_unique(rep, sheet, rows, key):
    seen = Counter(str(r[key]) for r in rows if r.get(key) is not None)
    for val, n in seen.items():
        if n > 1:
            rep.error(sheet, None, f"{key} '{val}' muncul {n} kali; harus unik")


def _check_ref(rep, sheet, rows, column, universe, universe_name, required=True):
    for row in rows:
        v = row.get(column)
        if v is None:
            if required:
                rep.error(sheet, row["_row"], f"{column} kosong")
            continue
        if str(v) not in universe:
            rep.error(
                sheet,
                row["_row"],
                f"{column}='{v}' tidak ada di {universe_name}",
            )


def _num(v):
    return v if isinstance(v, (int, float)) else None


def validate(path: Path) -> Report:
    rep = Report()
    wb = load_workbook(path, data_only=True)
    spec = _spec_map()

    missing = [s for s in spec if s not in wb.sheetnames]
    if missing:
        rep.error("<workbook>", None, f"sheet hilang: {', '.join(missing)}")
        return rep

    data = {}
    for name in spec:
        headers, rows = _read_sheet(wb[name])
        data[name] = rows
        for col in spec[name]:
            if col not in headers:
                rep.error(name, None, f"kolom '{col}' tidak ditemukan di header")

    if not rep.ok:
        return rep

    gi, bay, tmpl = data["01_GI"], data["02_BAY"], data["03_TEMPLATE"]
    samb, salur, trafo = data["04_SAMBUNGAN"], data["05_SALURAN"], data["06_TRAFO"]
    beban, gen = data["07_BEBAN"], data["08_PEMBANGKIT"]
    diameter, asumsi = data["10_DIAMETER"], data["09_ASUMSI"]
    skenario, aliran = data["11_SKENARIO"], data["12_ALIRAN"]
    tegangan, kapasitor = data["13_TEGANGAN"], data["14_KAPASITOR"]

    for name in spec:
        _check_required(rep, name, data[name], spec[name])

    # --- Identity uniqueness -------------------------------------------------
    _check_unique(rep, "01_GI", gi, "gi_id")
    _check_unique(rep, "02_BAY", bay, "bay_id")
    _check_unique(rep, "03_TEMPLATE", tmpl, "template_id")
    _check_unique(rep, "04_SAMBUNGAN", samb, "sambungan_id")
    _check_unique(rep, "10_DIAMETER", diameter, "diameter_id")
    _check_unique(rep, "11_SKENARIO", skenario, "skenario_id")
    _check_unique(rep, "12_ALIRAN", aliran, "aliran_id")
    _check_unique(rep, "06_TRAFO", trafo, "trafo_id")
    _check_unique(rep, "08_PEMBANGKIT", gen, "gen_id")
    _check_unique(rep, "09_ASUMSI", asumsi, "asumsi_id")

    gi_ids = {str(r["gi_id"]) for r in gi if r.get("gi_id")}
    bay_ids = {str(r["bay_id"]) for r in bay if r.get("bay_id")}
    tmpl_ids = {str(r["template_id"]) for r in tmpl if r.get("template_id")}
    samb_ids = {str(r["sambungan_id"]) for r in samb if r.get("sambungan_id")}

    # --- Controlled vocabularies --------------------------------------------
    _check_vocab(rep, "01_GI", gi, "skema_busbar", SKEMA_BUSBAR)
    _check_vocab(rep, "01_GI", gi, "busbar_sumber", SUMBER)
    _check_vocab(rep, "01_GI", gi, "busbar_confidence", CONFIDENCE)
    _check_vocab(rep, "02_BAY", bay, "bay_tipe", BAY_TIPE)
    _check_vocab(rep, "02_BAY", bay, "bus_normal", BUS_NORMAL)
    _check_vocab(rep, "04_SAMBUNGAN", samb, "phi", PHI)
    _check_vocab(rep, "07_BEBAN", beban, "titik_ukur", TITIK_UKUR)
    for name, rows in data.items():
        if any("sumber" in r for r in rows):
            _check_vocab(rep, name, rows, "sumber", SUMBER)
        if any("confidence" in r for r in rows):
            _check_vocab(rep, name, rows, "confidence", CONFIDENCE)
    _check_vocab(rep, "03_TEMPLATE", tmpl, "status_rekonstruksi", CONFIDENCE)

    # --- Cross-sheet references ---------------------------------------------
    _check_ref(rep, "02_BAY", bay, "gi_id", gi_ids, "01_GI")
    _check_ref(rep, "02_BAY", bay, "template_id", tmpl_ids, "03_TEMPLATE")
    _check_ref(rep, "04_SAMBUNGAN", samb, "dari_gi", gi_ids, "01_GI")
    _check_ref(rep, "04_SAMBUNGAN", samb, "ke_gi", gi_ids, "01_GI")
    _check_ref(rep, "04_SAMBUNGAN", samb, "dari_bay", bay_ids, "02_BAY")
    _check_ref(rep, "04_SAMBUNGAN", samb, "ke_bay", bay_ids, "02_BAY", required=False)
    _check_ref(rep, "05_SALURAN", salur, "sambungan_id", samb_ids, "04_SAMBUNGAN")
    _check_ref(rep, "10_DIAMETER", diameter, "gi_id", gi_ids, "01_GI")
    _check_ref(rep, "10_DIAMETER", diameter, "bay_id", bay_ids, "02_BAY")

    skn_ids = {str(r["skenario_id"]) for r in skenario if r.get("skenario_id")}
    ruas_ids = {str(r["ruas_id"]) for r in samb if r.get("ruas_id")}
    _check_ref(rep, "07_BEBAN", beban, "skenario_id", skn_ids, "11_SKENARIO")
    _check_ref(rep, "12_ALIRAN", aliran, "skenario_id", skn_ids, "11_SKENARIO")
    _check_ref(rep, "12_ALIRAN", aliran, "ruas_id", ruas_ids, "04_SAMBUNGAN")
    _check_ref(rep, "12_ALIRAN", aliran, "dari_gi", gi_ids, "01_GI")
    _check_ref(rep, "12_ALIRAN", aliran, "ke_gi", gi_ids, "01_GI")
    _check_vocab(rep, "07_BEBAN", beban, "jenis_nilai", JENIS_NILAI)
    _check_ref(rep, "06_TRAFO", trafo, "gi_id", gi_ids, "01_GI")
    _check_ref(rep, "06_TRAFO", trafo, "bay_id", bay_ids, "02_BAY")
    _check_ref(rep, "07_BEBAN", beban, "gi_id", gi_ids, "01_GI")
    _check_ref(rep, "08_PEMBANGKIT", gen, "gi_id", gi_ids, "01_GI")
    _check_ref(rep, "08_PEMBANGKIT", gen, "bay_id", bay_ids, "02_BAY")

    # --- Rule: bay type must agree with its template -------------------------
    tmpl_by_id = {str(r["template_id"]): r for r in tmpl if r.get("template_id")}
    expect = {
        "PENGHANTAR": "PHT",
        "TRAFO": "TRF",
        "GENERATOR": "GEN",
        "KOPEL": "KOPEL",
    }
    for row in bay:
        t, tid = row.get("bay_tipe"), row.get("template_id")
        if not t or not tid:
            continue
        token = expect.get(str(t))
        if token and token not in str(tid):
            rep.warn(
                "02_BAY",
                row["_row"],
                f"bay_tipe='{t}' tampak tidak cocok dengan template_id='{tid}'",
            )

    # --- Rule: bay template must match the GI's declared busbar scheme -------
    # Busbar arrangement belongs to the GI. A double-busbar template inside a
    # GI declared single busbar is a contradiction that must not reach CIM.
    gi_skema = {str(r["gi_id"]): str(r.get("skema_busbar")) for r in gi if r.get("gi_id")}
    for row in bay:
        g = str(row.get("gi_id"))
        tid = str(row.get("template_id"))
        t_sk = str(tmpl_by_id.get(tid, {}).get("skema_bus"))
        g_sk = gi_skema.get(g)
        if g_sk is None or t_sk == "None":
            continue
        if t_sk != g_sk:
            rep.error(
                "02_BAY",
                row["_row"],
                f"template '{tid}' berskema {t_sk} tetapi GI '{g}' dideklarasikan {g_sk} di 01_GI",
            )

    # --- Rule: bus_normal depends on the busbar scheme -----------------------
    # Only a real double busbar offers a choice of normal bus.
    for row in bay:
        tipe = str(row.get("bay_tipe"))
        bn = str(row.get("bus_normal"))
        g_sk = gi_skema.get(str(row.get("gi_id")))

        if tipe in ("KOPEL", "BUS_SECTION", "DIAMETER") and bn != "-":
            rep.error(
                "02_BAY",
                row["_row"],
                f"bay {tipe} harus bus_normal='-' (tidak memilih satu rel)",
            )
            continue

        if tipe in ("KOPEL", "BUS_SECTION", "DIAMETER"):
            continue

        if g_sk in ("DOUBLE", "DOUBLE_SECTION") and bn not in ("A", "B"):
            rep.error(
                "02_BAY",
                row["_row"],
                "bay di GI double busbar wajib bus_normal='A' atau 'B'",
            )
        if g_sk in ("SINGLE", "SINGLE_SECTION") and bn in ("A", "B"):
            rep.error(
                "02_BAY",
                row["_row"],
                f"bus_normal='{bn}' tidak berlaku: GI ini single busbar, tidak ada pilihan rel",
            )

    # --- Rule: one-and-a-half CB stations need diameter rows -----------------
    # On 1-1/2 CB the centre breaker is SHARED by two circuits, so the
    # one-bay-one-CB assumption behind every other template does not hold.
    dia_bays = {str(r.get("bay_id")) for r in diameter}
    for row in bay:
        if str(row.get("bay_tipe")) != "DIAMETER":
            continue
        bid = str(row.get("bay_id"))
        if bid not in dia_bays:
            rep.error(
                "02_BAY",
                row["_row"],
                f"bay DIAMETER '{bid}' belum punya baris di 10_DIAMETER",
            )

    for g, sk in gi_skema.items():
        if sk != "ONE_HALF_CB":
            continue
        n = sum(1 for r in bay if str(r.get("gi_id")) == g and str(r.get("bay_tipe")) == "DIAMETER")
        if n == 0:
            rep.error(
                "01_GI",
                None,
                f"GI '{g}' berskema ONE_HALF_CB tetapi tidak punya bay bertipe DIAMETER",
            )

    for row in diameter:
        g_sk = gi_skema.get(str(row.get("gi_id")))
        if g_sk != "ONE_HALF_CB":
            rep.error(
                "10_DIAMETER",
                row["_row"],
                f"GI '{row.get('gi_id')}' bukan ONE_HALF_CB; sheet ini hanya untuk 1-1/2 CB",
            )
        if str(row.get("cb_tengah_bersama")).upper() not in ("YA", "TIDAK"):
            rep.error("10_DIAMETER", row["_row"], "cb_tengah_bersama harus 'YA' atau 'TIDAK'")
        if not row.get("posisi_2"):
            rep.warn(
                "10_DIAMETER",
                row["_row"],
                "posisi_2 kosong: pasangan sirkit dalam diameter belum ditetapkan",
            )

    # --- Rule: phi must agree with the circuit count -------------------------
    by_ruas = defaultdict(list)
    for row in samb:
        by_ruas[str(row.get("ruas_id"))].append(row)

    expect_phi = {"SINGLE": 1, "DOUBLE": 2}
    for ruas, rows in by_ruas.items():
        phis = {str(r.get("phi")) for r in rows}
        if len(phis) > 1:
            rep.error("04_SAMBUNGAN", None, f"ruas '{ruas}' punya nilai phi berbeda: {phis}")
        counts = {r.get("jml_sirkit") for r in rows}
        if len(counts) > 1:
            rep.error("04_SAMBUNGAN", None, f"ruas '{ruas}' punya jml_sirkit berbeda: {counts}")

        phi = next(iter(phis))
        declared = next(iter(counts))
        want = expect_phi.get(phi)
        if want is not None and declared != want:
            rep.error(
                "04_SAMBUNGAN",
                rows[0]["_row"],
                f"phi='{phi}' mengharuskan jml_sirkit={want}, tertulis {declared}",
            )
        if phi == "MULTI" and isinstance(declared, (int, float)) and declared < 3:
            rep.error(
                "04_SAMBUNGAN",
                rows[0]["_row"],
                f"phi='MULTI' mengharuskan jml_sirkit>=3, tertulis {declared}",
            )

        # Every declared circuit should have its own row.
        if isinstance(declared, (int, float)) and len(rows) != declared:
            rep.warn(
                "04_SAMBUNGAN",
                None,
                f"ruas '{ruas}': jml_sirkit={int(declared)} tetapi hanya {len(rows)} baris sirkit terisi",
            )
        seq = sorted(r.get("sirkit_ke") for r in rows if r.get("sirkit_ke") is not None)
        if len(set(seq)) != len(seq):
            rep.error("04_SAMBUNGAN", None, f"ruas '{ruas}' punya sirkit_ke duplikat")

    # --- Rule: corridor flow must never become a load object -----------------
    # If a GI's load is already modelled, the flow on the line feeding it is a
    # consequence, not an extra consumption. Turning flow figures into
    # EnergyConsumer objects double-counts system load.
    flow_keys = {
        (str(r.get("skenario_id")), str(r.get("ke_gi")), _num(r.get("p_mw")))
        for r in aliran
    }
    for row in beban:
        key = (str(row.get("skenario_id")), str(row.get("gi_id")), _num(row.get("p_mw")))
        if key in flow_keys and key[2] is not None:
            rep.error(
                "07_BEBAN",
                row["_row"],
                f"p_mw={key[2]} untuk GI '{key[1]}' identik dengan aliran penghantar di "
                "12_ALIRAN pada skenario yang sama. Ini tanda beban dihitung dua kali: "
                "aliran penghantar adalah AKIBAT beban, bukan beban tambahan.",
            )

    # --- Rule: snapshot values need a scenario, ratings must not have one ----
    for row in beban:
        jn = str(row.get("jenis_nilai"))
        if jn == "RATING":
            rep.error(
                "07_BEBAN",
                row["_row"],
                "jenis_nilai=RATING tidak sah di sheet beban; rating trafo milik 06_TRAFO",
            )
        if jn == "DUMMY" and str(row.get("confidence")) != "ASSUMED":
            rep.error(
                "07_BEBAN",
                row["_row"],
                "jenis_nilai=DUMMY wajib confidence=ASSUMED (nilai uji, bukan data)",
            )

    # --- Rule: scenario metadata must be complete enough to be reproducible --
    for row in skenario:
        if not row.get("tanggal"):
            rep.error("11_SKENARIO", row["_row"], "tanggal wajib: snapshot tanpa waktu tidak bermakna")
        bs = _num(row.get("beban_sistem_mw"))
        if bs is not None and bs <= 0:
            rep.error("11_SKENARIO", row["_row"], "beban_sistem_mw harus > 0")

    unused_skn = skn_ids - {str(r.get("skenario_id")) for r in beban} - {
        str(r.get("skenario_id")) for r in aliran
    }
    if unused_skn:
        rep.warn("11_SKENARIO", None, f"skenario tidak dipakai: {', '.join(sorted(unused_skn))}")

    # --- Rule: measured voltages must be physically plausible ----------------
    _check_unique(rep, "13_TEGANGAN", tegangan, "tegangan_id")
    _check_unique(rep, "14_KAPASITOR", kapasitor, "kapasitor_id")
    _check_ref(rep, "13_TEGANGAN", tegangan, "gi_id", gi_ids, "01_GI")
    _check_ref(rep, "13_TEGANGAN", tegangan, "skenario_id", skn_ids, "11_SKENARIO")
    _check_ref(rep, "14_KAPASITOR", kapasitor, "gi_id", gi_ids, "01_GI")

    for row in tegangan:
        kv = _num(row.get("kv_terukur"))
        nom = _num(row.get("kv_nominal"))
        if kv is None or nom is None or nom <= 0:
            continue
        pu = kv / nom
        if not (0.85 <= pu <= 1.15):
            msg = (
                f"kv_terukur={kv} pada nominal {nom} kV = {pu:.2f} pu, di luar "
                "0,85-1,15. Periksa pembacaan angka sebelum dipakai."
            )
            if str(row.get("confidence")) == "UNKNOWN":
                rep.warn("13_TEGANGAN", row["_row"], msg + " (sudah ditandai UNKNOWN)")
            else:
                rep.error("13_TEGANGAN", row["_row"], msg)

    for row in kapasitor:
        m = _num(row.get("mvar_per_step"))
        if m is not None and m <= 0:
            rep.error("14_KAPASITOR", row["_row"], "mvar_per_step harus > 0")

    # --- Sanity: modelled load vs the scenario's stated system peak ----------
    # Not an error: the figure covers only part of the system and losses are
    # excluded. A large gap simply means coverage is still incomplete.
    for sc in skenario:
        sid = str(sc.get("skenario_id"))
        target = _num(sc.get("beban_sistem_mw"))
        if target is None:
            continue
        total = sum(
            _num(r.get("p_mw")) or 0.0 for r in beban if str(r.get("skenario_id")) == sid
        )
        if total > 0:
            pct = 100.0 * total / target
            rep.info.append(
                f"skenario {sid}: beban termodel {total:.1f} MW vs puncak sistem "
                f"{target:.0f} MW ({pct:.0f}%)"
            )
            if pct > 110:
                rep.error(
                    "07_BEBAN",
                    None,
                    f"total beban termodel ({total:.1f} MW) melebihi 110% beban puncak "
                    f"sistem ({target:.0f} MW) pada skenario {sid}. Periksa beban ganda.",
                )

    # --- Rule: power balance figures must be internally consistent ----------
    neraca = data["15_NERACA_DAYA"]
    _check_unique(rep, "15_NERACA_DAYA", neraca, "neraca_id")
    for row in neraca:
        parts = [_num(row.get(k)) for k in ("transfer_sklt_mw", "kit_non_bbm_mw", "kit_bbm_mw")]
        dmn = _num(row.get("dmn_total_mw"))
        if dmn is not None and all(v is not None for v in parts):
            total = sum(parts)
            if abs(total - dmn) > 0.5:
                rep.error(
                    "15_NERACA_DAYA",
                    row["_row"],
                    f"komponen pasokan berjumlah {total:.1f} MW tetapi dmn_total_mw={dmn}",
                )
        bp = _num(row.get("beban_puncak_mw"))
        cad = _num(row.get("cadangan_mw"))
        if dmn is not None and bp is not None and cad is not None:
            if abs((dmn - bp) - cad) > 0.5:
                rep.error(
                    "15_NERACA_DAYA",
                    row["_row"],
                    f"cadangan_mw={cad} tidak sama dengan DMN-beban ({dmn - bp:.1f} MW)",
                )
        if bp is not None and bp <= 0:
            rep.error("15_NERACA_DAYA", row["_row"], "beban_puncak_mw harus > 0")
        # A reserve smaller than the largest unit fails the stated N-1 criterion.
        unit = _num(row.get("unit_terbesar_mw"))
        if cad is not None and unit is not None and cad < unit:
            rep.info.append(
                f"neraca {row.get('neraca_id')}: cadangan {cad:.0f} MW < unit terbesar "
                f"{unit:.0f} MW -> kriteria N-1 tidak terpenuhi (sesuai catatan buku)"
            )

    # --- Rule: busbar claims must not be overstated --------------------------
    for row in gi:
        if str(row.get("busbar_confidence")) == "VERIFIED" and not row.get("busbar_dok_ref"):
            rep.error(
                "01_GI",
                row["_row"],
                "busbar_confidence=VERIFIED wajib menyebut busbar_dok_ref",
            )

    # --- Rule: zero is not 'unknown' -----------------------------------------
    for row in salur:
        for col in ("r_ohm", "x_ohm", "b_us", "panjang_km"):
            v = _num(row.get(col))
            if v is not None and v == 0 and str(row.get("confidence")) != "VERIFIED":
                rep.error(
                    "05_SALURAN",
                    row["_row"],
                    f"{col}=0 tanpa confidence=VERIFIED. Kosongkan sel bila nilainya belum diketahui.",
                )
        x = _num(row.get("x_ohm"))
        if x is not None and x < 0:
            rep.error("05_SALURAN", row["_row"], "x_ohm negatif")

    # --- Rule: transformer windings must differ ------------------------------
    for row in trafo:
        hv, lv = _num(row.get("hv_kv")), _num(row.get("lv_kv"))
        if hv is not None and lv is not None and hv == lv:
            rep.error(
                "06_TRAFO",
                row["_row"],
                f"hv_kv dan lv_kv sama ({hv}). Ini pola cacat yang ada di file aset PoC.",
            )
        if hv is not None and lv is not None and lv > hv:
            rep.warn("06_TRAFO", row["_row"], "lv_kv lebih besar dari hv_kv")
        mva = _num(row.get("mva"))
        if mva is not None and mva <= 0:
            rep.error("06_TRAFO", row["_row"], "mva harus > 0")

    # --- Rule: load sanity ---------------------------------------------------
    for row in beban:
        p, q = _num(row.get("p_mw")), _num(row.get("q_mvar"))
        cp = _num(row.get("cos_phi"))
        if p is None and q is None:
            rep.warn("07_BEBAN", row["_row"], "p_mw dan q_mvar keduanya kosong")
        if cp is not None and not (0 < cp <= 1):
            rep.error("07_BEBAN", row["_row"], f"cos_phi={cp} di luar rentang (0,1]")
        if str(row.get("titik_ukur")) == "TIDAK_DIKETAHUI":
            rep.warn(
                "07_BEBAN",
                row["_row"],
                "titik_ukur TIDAK_DIKETAHUI; nilai tidak bisa dipetakan ke sisi trafo",
            )

    # --- Rule: every ASSUMED value needs a register entry --------------------
    registered = defaultdict(set)
    for row in asumsi:
        registered[str(row.get("ref_sheet"))].add(str(row.get("ref_id")))

    id_col = {
        "02_BAY": "bay_id",
        "04_SAMBUNGAN": "sambungan_id",
        "05_SALURAN": "sambungan_id",
        "06_TRAFO": "trafo_id",
        "07_BEBAN": "beban_id",
        "08_PEMBANGKIT": "gen_id",
        "10_DIAMETER": "diameter_id",
    }
    for sheet, col in id_col.items():
        for row in data[sheet]:
            if str(row.get("confidence")) != "ASSUMED":
                continue
            rid = str(row.get(col))
            if rid not in registered.get(sheet, set()):
                rep.error(
                    sheet,
                    row["_row"],
                    f"confidence=ASSUMED tetapi '{rid}' tidak terdaftar di 09_ASUMSI",
                )

    # --- Rule: reconstructed templates must be flagged -----------------------
    for row in tmpl:
        if str(row.get("status_rekonstruksi")) == "VERIFIED" and not row.get("dasar_rekonstruksi"):
            rep.error("03_TEMPLATE", row["_row"], "status VERIFIED wajib menyebut dasar_rekonstruksi")

    unverified = [r for r in tmpl if str(r.get("status_rekonstruksi")) != "VERIFIED"]
    if unverified:
        rep.warn(
            "03_TEMPLATE",
            None,
            f"{len(unverified)} dari {len(tmpl)} template belum VERIFIED "
            "(wajar: template ini rekonstruksi, bukan standar PLN resmi)",
        )

    # --- Coverage summary ----------------------------------------------------
    used_tmpl = {str(r.get("template_id")) for r in bay}
    unused = tmpl_ids - used_tmpl
    if unused:
        rep.info.append(f"template belum dipakai: {', '.join(sorted(unused))}")

    bays_per_gi = Counter(str(r.get("gi_id")) for r in bay)
    for g in sorted(gi_ids):
        if bays_per_gi.get(g, 0) == 0:
            rep.warn("01_GI", None, f"GI '{g}' belum punya bay di 02_BAY")

    filled = sum(1 for r in salur if _num(r.get("x_ohm")) is not None)
    rep.info.append(f"impedansi saluran terisi: {filled}/{len(salur)}")
    filled_t = sum(1 for r in trafo if _num(r.get("z_persen")) is not None)
    rep.info.append(f"impedansi trafo terisi: {filled_t}/{len(trafo)}")
    rep.info.append(f"total bay: {len(bay)} di {len(gi_ids)} GI")

    skema_count = Counter(str(r.get("skema_busbar")) for r in gi)
    rep.info.append(
        "skema busbar: " + ", ".join(f"{k}={v}" for k, v in sorted(skema_count.items()))
    )
    verified_bb = sum(1 for r in gi if str(r.get("busbar_confidence")) == "VERIFIED")
    rep.info.append(f"skema busbar VERIFIED: {verified_bb}/{len(gi)} GI")

    phi_count = Counter(str(r.get("phi")) for r in samb)
    rep.info.append("phi ruas: " + ", ".join(f"{k}={v}" for k, v in sorted(phi_count.items())))

    if diameter:
        rep.info.append(f"diameter 1-1/2 CB: {len(diameter)} baris")
    if skenario:
        rep.info.append(f"skenario: {len(skenario)}; aliran penghantar tercatat: {len(aliran)}")

    return rep


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    path = Path(argv[1])
    if not path.exists():
        print(f"file tidak ditemukan: {path}")
        return 2

    rep = validate(path)

    print(f"=== VALIDASI {path.name} ===\n")

    if rep.errors:
        print(f"ERROR ({len(rep.errors)}) -- generator CIM tidak boleh dijalankan:")
        for e in rep.errors:
            print(f"  [X] {e}")
        print()

    if rep.warnings:
        print(f"PERINGATAN ({len(rep.warnings)}):")
        for w in rep.warnings:
            print(f"  [!] {w}")
        print()

    if rep.info:
        print("RINGKASAN:")
        for i in rep.info:
            print(f"  [i] {i}")
        print()

    print("HASIL:", "LULUS" if rep.ok else "GAGAL")
    return 0 if rep.ok else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
